"""
Run the Birdsong decoder over every wild Bewick's wren recording and write an
Excel report.

Point of the exercise: these clips are REAL birds — none of them carries a
message. So the decoded text column is expected to be nonsense throughout, and
`is_message` is expected to be FALSE for all 222. That makes this a false-positive
audit of the detector, not a transcription job: the interesting columns are
`viterbi_score` (how close each wild clip came to passing as a message) and the
margin below the threshold.

    ../.venv/bin/python decode_report.py
"""

import argparse
import glob
import os
import sys

import numpy as np
import soundfile as sf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.append(os.path.join(os.path.dirname(__file__), "..", "engine"))

import generate as G
import species as SP
from config import SAMPLE_RATE

COLUMNS = [
    ("clip", 22), ("duration_s", 11), ("peak_amp", 10), ("is_message", 12),
    ("viterbi_score", 14), ("margin_vs_threshold", 19), ("n_symbols", 11),
    ("decoded_text", 60),
]


def analyse(path):
    audio, sr = sf.read(path)
    audio = np.ascontiguousarray(audio, dtype=np.float32)
    if audio.ndim > 1:                       # fold any stereo down to mono
        audio = audio.mean(axis=1)
    ok, text, score = G.is_message(audio)
    return {
        "clip": os.path.basename(path),
        "duration_s": round(len(audio) / sr, 2),
        "peak_amp": round(float(np.abs(audio).max()), 3),
        "is_message": "TRUE" if ok else "FALSE",
        "viterbi_score": round(score, 4),
        "margin_vs_threshold": round(score - G._MESSAGE_THRESHOLD, 4),
        "n_symbols": len(text.replace(" ", "")),
        "decoded_text": text,
    }


def main():
    ap = argparse.ArgumentParser()
    SP.add_argument(ap)
    args = ap.parse_args()
    sp = SP.get(args.species)
    print(SP.banner(sp))

    out_path = os.path.join(os.path.dirname(__file__), "out",
                            f"{sp.code}_decode_report.xlsx")
    paths = sorted(glob.glob(os.path.join(sp.clips_dir, "*.wav")))
    if not paths:
        raise SystemExit(f"no .wav files found in {sp.clips_dir}")

    rows = []
    for i, p in enumerate(paths, 1):
        rows.append(analyse(p))
        if i % 25 == 0 or i == len(paths):
            print(f"  decoded {i}/{len(paths)}")

    wb = Workbook()
    ws = wb.active
    ws.title = "decode results"

    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF")
    for c, (name, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.fill, cell.font = header_fill, header_font
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = width

    flagged = PatternFill("solid", fgColor="F8CBAD")     # any clip that passed
    for r, row in enumerate(rows, 2):
        for c, (name, _) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=r, column=c, value=row[name])
            if name == "decoded_text":
                cell.font = Font(name="Menlo", size=10)
            if row["is_message"] == "TRUE":
                cell.fill = flagged
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    # A second sheet holding the reading of these numbers, so the file explains
    # itself to anyone who opens it without this conversation in front of them.
    scores = np.array([r["viterbi_score"] for r in rows])
    n_flagged = sum(r["is_message"] == "TRUE" for r in rows)
    s2 = wb.create_sheet("summary")
    for r, (k, v) in enumerate([
        ("species", sp.name),
        ("clips analysed", len(rows)),
        ("all are WILD recordings (no message encoded)", "yes"),
        ("", ""),
        ("flagged as a message (false positives)", n_flagged),
        ("false-positive rate", f"{100 * n_flagged / len(rows):.1f}%"),
        ("", ""),
        ("message threshold", G._MESSAGE_THRESHOLD),
        ("wild score — max", round(float(scores.max()), 4)),
        ("wild score — mean", round(float(scores.mean()), 4)),
        ("wild score — min", round(float(scores.min()), 4)),
        ("closest clip to threshold", rows[int(scores.argmax())]["clip"]),
        ("its margin below threshold", round(float(scores.max() - G._MESSAGE_THRESHOLD), 4)),
        ("", ""),
        ("decoded_text is EXPECTED to be nonsense", "these are real birds, not messages"),
        ("the decoder is closed-set: it always returns its best guess", ""),
        ("so is_message / viterbi_score decide whether to believe it", ""),
    ], 1):
        s2.cell(row=r, column=1, value=k).font = Font(bold=True)
        s2.cell(row=r, column=2, value=v)
    s2.column_dimensions["A"].width = 58
    s2.column_dimensions["B"].width = 38

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)

    print(f"\nwrote {out_path}")
    print(f"  {len(rows)} clips, {n_flagged} flagged as messages "
          f"({100 * n_flagged / len(rows):.1f}% false positive)")
    print(f"  wild scores: max {scores.max():+.4f}  mean {scores.mean():+.4f}  "
          f"(threshold {G._MESSAGE_THRESHOLD})")


if __name__ == "__main__":
    main()
